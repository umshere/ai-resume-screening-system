import asyncio
import streamlit as st
import pandas as pd
import io
import PyPDF2
import docx
import datetime
from typing import List, Dict
from dotenv import load_dotenv
import os

# Load environment variables from .env file
load_dotenv()

from semantic_kernel.contents import ChatMessageContent, AuthorRole
from src.mas import Orchestrator, MultiAgent
from src.usage_protection import UsageTracker, add_usage_monitoring, show_usage_stats

st.set_page_config(
    page_title="AI Resume Screening & Matching System",
    page_icon="📋",
    layout="wide"
)

# Apple-inspired minimal CSS design
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif');
    
    /* Clean background and typography */
    .main {
        background: #f8f9fa;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    }
    
    /* Minimal step indicator */
    .step-indicator {
        display: flex;
        justify-content: space-between;
        align-items: center;
        background: white;
        border-radius: 12px;
        padding: 16px 24px;
        margin: 24px 0;
        box-shadow: 0 1px 3px rgba(0,0,0,0.08);
        border: 1px solid #e5e7eb;
    }
    
    .step {
        display: flex;
        align-items: center;
        gap: 8px;
        font-size: 14px;
        font-weight: 500;
        color: #6b7280;
    }
    
    .step.completed {
        color: #059669;
    }
    
    .step-circle {
        width: 24px;
        height: 24px;
        border-radius: 50%;
        background: #e5e7eb;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 12px;
        font-weight: 600;
        color: #6b7280;
    }
    
    .step-circle.completed {
        background: #059669;
        color: white;
    }
    
    .step-divider {
        flex: 1;
        height: 2px;
        background: #e5e7eb;
        margin: 0 16px;
    }
    
    .step-divider.completed {
        background: #059669;
    }
    
    /* Clean file uploader */
    .stFileUploader > div {
        background: white;
        border: 2px dashed #d1d5db;
        border-radius: 12px;
        padding: 32px;
        text-align: center;
        transition: all 0.2s ease;
    }
    
    .stFileUploader > div:hover {
        border-color: #3b82f6;
        background: #f8faff;
    }
    
    /* Drag and drop text area */
    .drag-drop-area {
        background: white;
        border: 2px dashed #d1d5db;
        border-radius: 12px;
        padding: 32px;
        text-align: center;
        transition: all 0.2s ease;
        cursor: pointer;
        margin: 16px 0;
    }
    
    .drag-drop-area:hover {
        border-color: #3b82f6;
        background: #f8faff;
    }
    
    .drag-drop-area.dragover {
        border-color: #059669;
        background: #f0fdf4;
    }
    
    /* Agent preview cards */
    .agent-preview-card {
        background: white;
        border: 1px solid #e5e7eb;
        border-radius: 12px;
        padding: 16px;
        margin: 8px 0;
        transition: all 0.2s ease;
        border-left: 4px solid #3b82f6;
    }
    
    .agent-preview-card:hover {
        box-shadow: 0 4px 12px rgba(0,0,0,0.1);
        transform: translateY(-2px);
    }
    
    .agent-title {
        font-weight: 600;
        color: #1f2937;
        margin-bottom: 4px;
    }
    
    .agent-description {
        color: #6b7280;
        font-size: 14px;
        line-height: 1.4;
    }
    
    /* Instructions banner */
    .instruction-banner {
        background: linear-gradient(135deg, #3b82f6, #1d4ed8);
        color: white;
        padding: 12px 20px;
        border-radius: 10px;
        margin-bottom: 20px;
        font-weight: 500;
        text-align: center;
        box-shadow: 0 2px 8px rgba(59, 130, 246, 0.2);
    }
    
    /* Clean buttons */
    .stButton > button {
        background: #3b82f6;
        color: white;
        border: none;
        border-radius: 8px;
        padding: 12px 24px;
        font-weight: 500;
        font-size: 14px;
        transition: all 0.2s ease;
        box-shadow: 0 1px 2px rgba(0,0,0,0.05);
    }
    
    .stButton > button:hover {
        background: #2563eb;
        transform: translateY(-1px);
        box-shadow: 0 4px 8px rgba(0,0,0,0.12);
    }
    
    /* Clean text areas */
    .stTextArea textarea {
        border: 1px solid #d1d5db;
        border-radius: 8px;
        font-size: 14px;
        line-height: 1.5;
    }
    
    .stTextArea textarea:focus {
        border-color: #3b82f6;
        box-shadow: 0 0 0 3px rgba(59, 130, 246, 0.1);
    }
    
    /* Clean tabs */
    .stTabs [data-baseweb="tab-list"] {
        gap: 4px;
        background: #f3f4f6;
        border-radius: 10px;
        padding: 4px;
    }
    
    .stTabs [data-baseweb="tab"] {
        background: transparent;
        border: none;
        border-radius: 6px;
        padding: 8px 16px;
        font-weight: 500;
        color: #6b7280;
        transition: all 0.2s ease;
    }
    
    .stTabs [aria-selected="true"] {
        background: white;
        color: #374151;
        box-shadow: 0 1px 2px rgba(0,0,0,0.05);
    }
    
    /* Clean agent status */
    .agent-status {
        background: white;
        border: 1px solid #e5e7eb;
        border-radius: 8px;
        padding: 12px 16px;
        margin: 8px 0;
        border-left: 3px solid #3b82f6;
    }
    
    /* Clean headers */
    h1, h2, h3 {
        color: #111827;
        font-weight: 600;
        letter-spacing: -0.025em;
    }
    
    /* Hide streamlit branding */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    /* Clean containers */
    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }
    </style>
    """,
    unsafe_allow_html=True
)

def extract_text_from_pdf(pdf_file) -> str:
    """Extract text from PDF file"""
    try:
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        st.error(f"Error reading PDF: {str(e)}")
        return ""

def extract_text_from_docx(docx_file) -> str:
    """Extract text from DOCX file"""
    try:
        doc = docx.Document(docx_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        st.error(f"Error reading DOCX: {str(e)}")
        return ""

def process_resume_files(uploaded_files) -> List[Dict]:
    """Process uploaded resume files and extract text"""
    resumes = []
    for uploaded_file in uploaded_files:
        st.write(f"🔍 Processing: {uploaded_file.name} (Type: {uploaded_file.type})")
        
        if uploaded_file.type == "application/pdf":
            text = extract_text_from_pdf(uploaded_file)
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            text = extract_text_from_docx(uploaded_file)
        elif uploaded_file.type == "text/plain":
            text = str(uploaded_file.read(), "utf-8")
        else:
            st.warning(f"Unsupported file type: {uploaded_file.type}")
            continue
            
        if text.strip():
            resumes.append({
                "filename": uploaded_file.name,
                "content": text
            })
            st.success(f"✅ Successfully extracted {len(text)} characters from {uploaded_file.name}")
        else:
            st.error(f"❌ Failed to extract text from {uploaded_file.name} - file appears to be empty or corrupted")
    
    st.write(f"📊 Total processed resumes: {len(resumes)}")
    return resumes

async def run_resume_screening(job_profile, resumes, num_agents):
    """Run the resume screening process"""
    # Create screening context
    screening_context = {
        "job_profile": job_profile,
        "resumes": resumes,
        "num_resumes": len(resumes)
    }
    
    orchestrator = Orchestrator(screening_context, num_agents)
    dynamic_agents = orchestrator.run()

    mas = MultiAgent()
    expert_agents = mas.create_agents(dynamic_agents)
    expert_agents_names = [agent.name for agent in expert_agents]

    return expert_agents, expert_agents_names, mas

async def main_screening(job_profile, resumes, num_agents, max_interactions=None):
    """Main screening process"""
    expert_agents, expert_agents_names, mas = await run_resume_screening(job_profile, resumes, num_agents)
    
    # Set max_interactions based on number of agents if not specified (optimized for speed)
    if max_interactions is None:
        max_interactions = max(num_agents + 1, 3)  # Just 1 round per agent + 1 final round, minimum 3
    
    with st.sidebar:
        st.title("🤖 Expert Screening Agents")
        agent_placeholders = {name: st.empty() for name in expert_agents_names}
        agent_status = {name: "⏳ Initializing..." for name in expert_agents_names}
        
        # Define agent specializations for better user understanding
        agent_specializations = {
            "Skills_Analysis_Agent": "🔧 Technical Skills & Expertise",
            "Experience_Evaluation_Agent": "📈 Work Experience & Career",
            "Education_Assessment_Agent": "🎓 Educational Background",
            "Cultural_Fit_Agent": "🤝 Team & Cultural Alignment",
            "Leadership_Assessment_Agent": "👑 Leadership & Management",
            "Technical_Depth_Agent": "🔬 Deep Technical Analysis"
        }
        
        for agent_name in expert_agents_names:
            specialization = agent_specializations.get(agent_name, "🔍 General Analysis")
            agent_placeholders[agent_name].markdown(
                f'<div class="agent-status">👤 <strong>{agent_name}</strong><br><small>{specialization}</small><br>⏳ Initializing...</div>', 
                unsafe_allow_html=True
            )
            await asyncio.sleep(1)

    selection_function = mas.create_selection_function(expert_agents)
    termination_keyword = 'yes'
    termination_function = mas.create_termination_function(termination_keyword)

    group = mas.create_chat_group(
        expert_agents,
        selection_function,
        termination_function,
        termination_keyword
    )

    interactions: int = 0
    is_complete: bool = False
    
    # Define detailed agent activities based on their roles
    def get_agent_activity(agent_name, round_num, total_resumes):
        """Get detailed activity description for each agent"""
        activities = {
            "Skills_Analysis_Agent": [
                f"🔍 Analyzing technical skills alignment across {total_resumes} resume(s)",
                f"⚙️ Evaluating programming languages and frameworks",
                f"🛠️ Assessing tool proficiency and certifications",
                f"📊 Cross-referencing skill requirements with candidate experience"
            ],
            "Experience_Evaluation_Agent": [
                f"📈 Evaluating work experience relevance for {total_resumes} candidate(s)",
                f"🏢 Analyzing company backgrounds and industry experience",
                f"📅 Assessing career progression and tenure patterns",
                f"🎯 Matching role responsibilities with job requirements"
            ],
            "Education_Assessment_Agent": [
                f"🎓 Reviewing educational qualifications for {total_resumes} resume(s)",
                f"🏫 Evaluating degree relevance and institution reputation",
                f"📚 Assessing additional certifications and training",
                f"🧠 Analyzing academic achievements and projects"
            ],
            "Cultural_Fit_Agent": [
                f"🤝 Analyzing cultural alignment indicators across {total_resumes} profile(s)",
                f"💭 Evaluating communication style and collaboration signals",
                f"🌟 Assessing leadership potential and team dynamics",
                f"🎯 Reviewing personality traits and work style indicators"
            ],
            "Leadership_Assessment_Agent": [
                f"👑 Evaluating leadership experience in {total_resumes} candidate(s)",
                f"📊 Analyzing team management and project leadership",
                f"🚀 Assessing strategic thinking and decision-making",
                f"🌟 Reviewing mentoring and development capabilities"
            ],
            "Technical_Depth_Agent": [
                f"🔬 Deep-diving into technical expertise across {total_resumes} resume(s)",
                f"⚡ Analyzing system architecture and design experience",
                f"🧪 Evaluating problem-solving and debugging skills",
                f"🔧 Assessing code quality and best practices knowledge"
            ]
        }
        
        # Get activities for this agent, default to generic if not found
        agent_activities = activities.get(agent_name, [
            f"🔍 Analyzing candidate profiles for {total_resumes} resume(s)",
            f"📊 Evaluating qualifications and requirements",
            f"🎯 Assessing job-candidate alignment",
            f"📈 Generating detailed analysis insights"
        ])
        
        # Return activity based on round number (cycle through activities)
        activity_index = (round_num - 1) % len(agent_activities)
        return agent_activities[activity_index]
    
    # Create screening context with recruiter preferences
    recruiter_priorities = st.session_state.get('recruiter_priorities', '')
    special_requirements = st.session_state.get('special_requirements', '')
    
    screening_input = f"""
    Job Profile: {job_profile}
    
    Number of Resumes to Screen: {len(resumes)}
    
    Recruiter Priorities: {recruiter_priorities if recruiter_priorities else 'None specified'}
    
    Special Requirements: {special_requirements if special_requirements else 'None specified'}
    
    Resumes:
    {chr(10).join([f"{i+1}. {resume['filename']}: {resume['content'][:500]}..." for i, resume in enumerate(resumes)])}
    
    Please analyze each resume against the job profile, taking into account the recruiter's priorities and special requirements. Provide matching scores with detailed explanations, highlighting how well each candidate meets the specific priorities and requirements mentioned.
    """
    
    progress_bar = st.progress(0)
    status_text = st.empty()
    detailed_status = st.empty()
    
    with st.spinner("🔍 Screening resumes..."):
        while not is_complete and interactions < max_interactions:
            # Update agent status with detailed activity
            current_agent = expert_agents_names[interactions % len(expert_agents_names)]
            current_round = interactions + 1
            
            # Get detailed activity for this agent
            current_activity = get_agent_activity(current_agent, current_round, len(resumes))
            
            agent_status[current_agent] = f"🔍 {current_activity.split(' ', 1)[1]}"  # Remove emoji from activity
            specialization = agent_specializations.get(current_agent, "🔍 General Analysis")
            agent_placeholders[current_agent].markdown(
                f'<div class="agent-status">👤 <strong>{current_agent}</strong><br><small>{specialization}</small><br>🔍 {current_activity.split(" ", 1)[1]}</div>', 
                unsafe_allow_html=True
            )
            
            # Update main status with detailed information
            phase = "⚡ Quick Analysis" if current_round <= max_interactions // 2 else "🎯 Final Scoring"
            status_text.markdown(f"**Round {current_round}/{max_interactions}** | **{current_agent}** | {len(expert_agents)} agents | *{phase}*")
            detailed_status.info(f"🎯 {current_activity}")
            progress_bar.progress(current_round / max_interactions)
            
            await group.add_chat_message(ChatMessageContent(role=AuthorRole.USER, content=screening_input))

            async for response in group.invoke_stream():
                # Check if the response contains termination keyword
                if termination_keyword.lower() in response.content.lower():
                    is_complete = True
                    break

            interactions += 1
            await asyncio.sleep(1)
    
    # Update final status with completion summary
    for agent_name in expert_agents_names:
        agent_status[agent_name] = "✅ Analysis Complete - Ready for results compilation"
        specialization = agent_specializations.get(agent_name, "🔍 General Analysis")
        agent_placeholders[agent_name].markdown(
            f'<div class="agent-status">👤 <strong>{agent_name}</strong><br><small>{specialization}</small><br>✅ Analysis Complete</div>', 
            unsafe_allow_html=True
        )
    
    progress_bar.progress(1.0)
    status_text.markdown("**✅ Resume screening completed!**")
    detailed_status.success(f"🎉 All {len(expert_agents)} agents have completed their analysis of {len(resumes)} resume(s). Compiling final results...")
# Streamlit UI for Resume Screening
def app():
    # Initialize session state variables
    if 'job_profile' not in st.session_state:
        st.session_state.job_profile = ""
    if 'resumes' not in st.session_state:
        st.session_state.resumes = []
    if 'num_agents' not in st.session_state:
        st.session_state.num_agents = 4
    if 'analysis_depth' not in st.session_state:
        st.session_state.analysis_depth = "Standard Analysis"
    if 'running' not in st.session_state:
        st.session_state.running = False
    if 'screening_results' not in st.session_state:
        st.session_state.screening_results = None
    if 'recruiter_priorities' not in st.session_state:
        st.session_state.recruiter_priorities = ""
    if 'special_requirements' not in st.session_state:
        st.session_state.special_requirements = ""

    st.title("AI Resume Screening")
    st.markdown("Intelligent resume screening with AI expert agents")
    
    # Usage monitoring (collapsed by default for cleaner UI)
    with st.expander("💰 Usage & Cost Monitoring"):
        show_usage_stats()
        st.caption("Daily limit: $10 | Session limit: 20 resumes | Real-time tracking")
    
    # Single minimal step indicator - Highlight Configuration first
    steps = ["Configuration", "Job Profile", "Resumes", "Review"]
    completed_steps = []
    
    # Configuration is always highlighted first as suggested
    completed_steps.append("Configuration")
    if st.session_state.job_profile.strip():
        completed_steps.append("Job Profile")
    if st.session_state.resumes:
        completed_steps.append("Resumes")
    if len(completed_steps) == 3:
        completed_steps.append("Review")
    
    # Clean step indicator
    step_html = '<div class="step-indicator">'
    for i, step in enumerate(steps):
        is_completed = step in completed_steps
        step_class = "step completed" if is_completed else "step"
        circle_class = "step-circle completed" if is_completed else "step-circle"
        
        step_html += f'''
            <div class="{step_class}">
                <div class="{circle_class}">{"✓" if is_completed else str(i+1)}</div>
                <span>{step}</span>
            </div>
        '''
        
        if i < len(steps) - 1:
            divider_class = "step-divider completed" if is_completed and steps[i+1] in completed_steps else "step-divider"
            step_html += f'<div class="{divider_class}"></div>'
    
    step_html += '</div>'
    st.markdown(step_html, unsafe_allow_html=True)
    
    # Clean tabs without duplicate status indicators - reordered with Configuration first
    tab1, tab2, tab3, tab4 = st.tabs(["Configuration", "Job Profile", "Resumes", "Review"])

    with tab1:
        # Configuration section - now first tab
        st.markdown('<div class="instruction-banner">🎯 Start here: Configure your AI screening agents and analysis preferences</div>', unsafe_allow_html=True)
        
        st.markdown("### ⚙️ Screening Configuration")
        st.markdown("Configure the AI agents and analysis depth for optimal screening results.")
        
        col_config1, col_config2, col_config3 = st.columns([1, 1, 1])
        
        with col_config1:
            st.markdown("**👥 Expert Agents**")
            num_agents = st.number_input(
                "Number of expert agents",
                min_value=2,
                max_value=6,
                value=3,  # Default to 3 for speed
                step=1,
                help="💡 More agents provide more detailed analysis but take longer to complete"
            )
            
        with col_config2:
            st.markdown("**🔍 Analysis Depth**")
            analysis_depth = st.selectbox(
                "Analysis depth",
                ["Quick Overview", "Standard Analysis", "Deep Dive"],
                index=0,  # Default to Quick for speed
                help="Choose analysis depth vs speed trade-off"
            )
            
        with col_config3:
            st.markdown("**⚡ Speed Mode**")
            speed_mode = st.checkbox(
                "Enable Speed Mode",
                value=True,
                help="Faster screening with optimized settings"
            )
            
            if speed_mode:
                st.info("⚡ Speed optimizations active")
                if num_agents > 3:
                    num_agents = 3
                    st.caption("→ Limited to 3 agents for speed")
                
        # Store in session state
        st.session_state.num_agents = num_agents
        st.session_state.analysis_depth = analysis_depth
        st.session_state.speed_mode = speed_mode
            
        # Agent explanation
        agent_descriptions = {
            2: "Basic analysis with 2 core agents",
            3: "Balanced analysis with 3 specialized agents", 
            4: "Comprehensive analysis with 4 expert agents (Recommended)",
            5: "Detailed analysis with 5 specialized agents",
            6: "Maximum analysis with 6 expert agents"
        }
        st.info(f"🤖 {agent_descriptions.get(num_agents, 'Custom configuration')}")
        
        # Continue with existing analysis depth code
        st.markdown("**🔍 Analysis Depth (Applied to all agents)**")
        analysis_depth = st.selectbox(
            "Analysis depth",
            ["Quick Overview", "Standard Analysis", "Deep Dive"],
            index=["Quick Overview", "Standard Analysis", "Deep Dive"].index(st.session_state.analysis_depth),
            help="Choose the level of detail for resume analysis"
        )
        st.session_state.analysis_depth = analysis_depth
        
        # Depth explanation
        depth_descriptions = {
            "Quick Overview": "⚡ Fast screening with key highlights",
            "Standard Analysis": "⚖️ Balanced analysis with detailed insights", 
            "Deep Dive": "🔬 Comprehensive analysis with detailed explanations"
        }
        st.info(depth_descriptions[analysis_depth])

        # Dynamic Agent Preview Section
        st.markdown("---")
        st.markdown("### 🤖 Your AI Screening Team")
        if num_agents != st.session_state.num_agents:
            st.info(f"🔄 Agent configuration updated: {st.session_state.num_agents} → {num_agents} agents")
        st.markdown(f"**{num_agents} Expert Agents** will be created for your screening process:")
        
        # Define agent types and responsibilities based on count
        def get_agent_lineup(count):
            base_agents = [
                {
                    "name": "Skills Analysis Agent",
                    "icon": "🔧",
                    "specialty": "Technical Skills & Expertise",
                    "duties": [
                        "Analyze technical skills alignment",
                        "Evaluate programming languages & frameworks", 
                        "Assess tool proficiency & certifications",
                        "Cross-reference requirements with experience"
                    ]
                },
                {
                    "name": "Experience Evaluation Agent", 
                    "icon": "📈",
                    "specialty": "Work Experience & Career",
                    "duties": [
                        "Evaluate work experience relevance",
                        "Analyze company backgrounds & industries",
                        "Assess career progression patterns", 
                        "Match role responsibilities with requirements"
                    ]
                },
                {
                    "name": "Education Assessment Agent",
                    "icon": "🎓", 
                    "specialty": "Educational Background",
                    "duties": [
                        "Review educational qualifications",
                        "Evaluate degree relevance & institutions",
                        "Assess certifications & additional training",
                        "Analyze academic achievements & projects"
                    ]
                },
                {
                    "name": "Cultural Fit Agent",
                    "icon": "🤝",
                    "specialty": "Team & Cultural Alignment", 
                    "duties": [
                        "Analyze cultural alignment indicators",
                        "Evaluate communication & collaboration signals",
                        "Assess personality traits & work style",
                        "Review team dynamics potential"
                    ]
                },
                {
                    "name": "Leadership Assessment Agent",
                    "icon": "👑",
                    "specialty": "Leadership & Management",
                    "duties": [
                        "Evaluate leadership experience",
                        "Analyze team management & project leadership", 
                        "Assess strategic thinking capabilities",
                        "Review mentoring & development skills"
                    ]
                },
                {
                    "name": "Technical Depth Agent", 
                    "icon": "🔬",
                    "specialty": "Deep Technical Analysis",
                    "duties": [
                        "Deep-dive into technical expertise",
                        "Analyze system architecture experience",
                        "Evaluate problem-solving & debugging skills", 
                        "Assess code quality & best practices knowledge"
                    ]
                }
            ]
            return base_agents[:count]
        
        agent_lineup = get_agent_lineup(num_agents)
        
        # Display agent cards in a responsive grid
        for i in range(0, len(agent_lineup), 2):
            cols = st.columns(2)
            for j, col in enumerate(cols):
                if i + j < len(agent_lineup):
                    agent = agent_lineup[i + j]
                    with col:
                        st.markdown(f'''
                        <div class="agent-preview-card">
                            <div class="agent-title">{agent["icon"]} {agent["name"]}</div>
                            <div class="agent-description">{agent["specialty"]}</div>
                        </div>
                        ''', unsafe_allow_html=True)
                        
                        with st.expander(f"View {agent['name']} duties", expanded=False):
                            st.markdown("**Key Responsibilities:**")
                            for duty in agent["duties"]:
                                st.write(f"• {duty}")
        
        # Recruiter Preferences Section
        st.markdown("---")
        st.markdown("### 👤 Recruiter Preferences")
        st.markdown("Add specific priorities or requirements to customize the screening process:")
        
        col_pref1, col_pref2 = st.columns([1, 1])
        
        with col_pref1:
            recruiter_priorities = st.text_area(
                "🎯 Hiring Priorities & Focus Areas",
                value=st.session_state.recruiter_priorities,
                height=100,
                placeholder="e.g., Prioritize candidates with startup experience, remote work capability, leadership potential...",
                help="Specify what matters most for this role"
            )
            st.session_state.recruiter_priorities = recruiter_priorities
            
        with col_pref2:
            special_requirements = st.text_area(
                "📋 Special Requirements & Deal-breakers",
                value=st.session_state.special_requirements, 
                height=100,
                placeholder="e.g., Must have security clearance, relocation not possible, specific visa requirements...",
                help="List any non-negotiable requirements"
            )
            st.session_state.special_requirements = special_requirements
        
        # Action buttons
        col_action1, col_action2, col_action3 = st.columns([1, 1, 1])
        with col_action2:
            if st.button("� Save Configuration", type="primary", use_container_width=True):
                st.success("✅ Configuration saved successfully!")
                st.info("📋 Next: Define your job profile in the Job Profile tab")

    with tab2:
        # Job Profile section
        st.markdown('<div class="instruction-banner">� Next: Define the job profile with requirements and qualifications</div>', unsafe_allow_html=True)
        
        st.markdown("### Job Description")
        
        # Drag and drop or text area for job profile
        job_input_method = st.radio(
            "Choose input method:",
            ["✏️ Type/Paste Job Description", "📁 Upload Job Description File"],
            horizontal=True
        )
        
        if job_input_method == "✏️ Type/Paste Job Description":
            job_profile_input = st.text_area(
                "Enter the job description or requirements:",
                value=st.session_state.job_profile,
                height=200,
                placeholder="Senior Software Engineer position requiring 5+ years experience in Python, React, and cloud technologies...",
                help="Include specific skills, experience requirements, and qualifications for better matching"
            )
        else:
            # File upload for job description
            st.markdown('''
            <div class="drag-drop-area">
                📁 Drag and drop job description file here<br>
                <small>Supported formats: TXT, PDF, DOCX</small>
            </div>
            ''', unsafe_allow_html=True)
            
            uploaded_job_file = st.file_uploader(
                "Or click to upload job description file:",
                type=['txt', 'pdf', 'docx'],
                help="Upload a file containing the job description"
            )
            
            job_profile_input = st.session_state.job_profile
            
            if uploaded_job_file:
                with st.spinner("📖 Processing job description file..."):
                    if uploaded_job_file.type == "text/plain":
                        job_content = str(uploaded_job_file.read(), "utf-8")
                    elif uploaded_job_file.type == "application/pdf":
                        job_content = extract_text_from_pdf(uploaded_job_file)
                    elif uploaded_job_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        job_content = extract_text_from_docx(uploaded_job_file)
                    else:
                        st.error("Unsupported file type")
                        job_content = ""
                    
                    if job_content.strip():
                        job_profile_input = job_content
                        st.success(f"✅ Successfully loaded job description ({len(job_content)} characters)")
                    else:
                        st.error("❌ Failed to extract content from file")
        
        st.session_state.job_profile = job_profile_input
        
        # Simple character counter
        char_count = len(job_profile_input)
        if char_count > 0:
            status_color = "🟢" if char_count >= 100 else "🟡"
            st.caption(f"{status_color} {char_count} characters")
        
        # Action button
        col_action = st.columns([2, 1, 2])
        with col_action[1]:
            if st.button("💾 Save Job Profile", type="primary", use_container_width=True, disabled=not job_profile_input.strip()):
                st.success("✅ Job profile saved!")
                st.info("📄 Next: Upload resumes in the Resumes tab")


    with tab3:
        # Resumes section with improved drag and drop
        st.markdown('<div class="instruction-banner">📄 Upload candidate resumes for AI-powered screening and analysis</div>', unsafe_allow_html=True)
        
        st.markdown("### 📄 Resumes")
        
        # Enhanced drag and drop area
        st.markdown('''
        <div class="drag-drop-area">
            📁 Drag and drop resume files here<br>
            <small>Supported formats: PDF, DOCX, TXT • Multiple files supported</small>
        </div>
        ''', unsafe_allow_html=True)
        
        uploaded_files_tab3 = st.file_uploader(
            "Or click to upload resume files:",
            type=['pdf', 'docx', 'txt'],
            accept_multiple_files=True,
            help="Support for PDF, DOCX, and TXT files. You can upload multiple resumes at once.",
            key="resume_uploader_tab3"
        )
        
        if uploaded_files_tab3:
            with st.spinner("📖 Processing uploaded resumes..."):
                st.info(f"🔍 Processing {len(uploaded_files_tab3)} uploaded file(s)...")
                newly_processed_resumes = process_resume_files(uploaded_files_tab3)
                st.success(f"✅ Successfully processed {len(newly_processed_resumes)} resume(s) from upload")
                
                # Append new resumes to existing ones in session state, avoiding duplicates by filename
                existing_filenames = {r['filename'] for r in st.session_state.resumes}
                added_count = 0
                for res in newly_processed_resumes:
                    if res['filename'] not in existing_filenames:
                        st.session_state.resumes.append(res)
                        existing_filenames.add(res['filename'])
                        added_count += 1
                    else:
                        st.warning(f"⚠️ Skipped duplicate file: {res['filename']}")
                
                if added_count > 0:
                    st.success(f"➕ Added {added_count} new resume(s) to your collection")
            
        if st.session_state.resumes:
            st.success(f"✅ Successfully processed {len(st.session_state.resumes)} resume(s)")
            
            # Action buttons
            col_btn1, col_btn2, col_btn3 = st.columns([1, 1, 1])
            with col_btn1:
                if st.button("🗑️ Clear All Resumes", help="Remove all uploaded files and start over", key="clear_resumes_button"):
                    st.session_state.resumes = []
                    st.rerun()
            with col_btn3:
                if st.button("� Save Resumes", type="primary", help="Confirm resume selection", key="save_resumes_button"):
                    st.success("✅ Resumes saved successfully!")
                    st.info("🚀 Next: Review and start screening in the Review tab")
            
            # Enhanced resume preview with file info
            with st.expander(f"📋 Preview Resumes ({len(st.session_state.resumes)})", expanded=False):
                for i, resume in enumerate(st.session_state.resumes):
                    # File info header
                    file_extension = resume['filename'].split('.')[-1].upper()
                    file_size = f"{len(resume['content'])} chars"
                    
                    col_info, col_remove = st.columns([4, 1])
                    with col_info:
                        st.markdown(f"**📄 {i+1}. {resume['filename']}** (`{file_extension}` • {file_size})")
                    with col_remove:
                        if st.button("❌", key=f"remove_{i}", help=f"Remove {resume['filename']}"):
                            st.session_state.resumes.pop(i)
                            st.rerun()
                    
                    # Content preview
                    preview_text = resume['content'][:300] + "..." if len(resume['content']) > 300 else resume['content']
                    st.text_area(
                        f"Content preview:",
                        value=preview_text,
                        height=100,
                        disabled=True,
                        key=f"preview_{i}"
                    )
                    if i < len(st.session_state.resumes) - 1:
                        st.divider()
        else:
            st.info("📤 Upload resume files to begin the screening process")

    with tab4:
        # Review inputs and final Start Screening button
        st.markdown('<div class="instruction-banner">🚀 Review your configuration and start the AI-powered resume screening</div>', unsafe_allow_html=True)
        
        st.markdown("### � Review & Start Screening")
        
        job_profile_tab4 = st.session_state.get('job_profile', "")
        resumes_tab4 = st.session_state.get('resumes', [])
        can_start = job_profile_tab4.strip() and resumes_tab4
        
        # Configuration Summary
        st.markdown("#### ⚙️ Configuration Summary")
        col_summary1, col_summary2 = st.columns([1, 1])
        
        with col_summary1:
            st.info(f"""
            **🤖 AI Agents:** {st.session_state.num_agents} expert agents
            **🔍 Analysis:** {st.session_state.analysis_depth}
            **📄 Resumes:** {len(resumes_tab4)} files uploaded
            """)
            
        with col_summary2:
            priorities_text = st.session_state.get('recruiter_priorities', 'None specified')[:50] + "..." if len(st.session_state.get('recruiter_priorities', '')) > 50 else st.session_state.get('recruiter_priorities', 'None specified')
            requirements_text = st.session_state.get('special_requirements', 'None specified')[:50] + "..." if len(st.session_state.get('special_requirements', '')) > 50 else st.session_state.get('special_requirements', 'None specified')
            
            st.info(f"""
            **🎯 Priorities:** {priorities_text}
            **� Requirements:** {requirements_text}
            **⏱️ Est. Time:** {max(st.session_state.num_agents // 2, 1)} minutes
            """)
        
        # Display detailed summaries with expandable sections
        if job_profile_tab4.strip():
            with st.expander("📝 Review Job Profile", expanded=False):
                st.text_area("Job Profile:", value=job_profile_tab4, height=150, disabled=True, key="review_job_profile")
        else:
            st.error("❌ Please provide a job profile in the 'Job Profile' tab.")

        if resumes_tab4:
            with st.expander(f"📄 Review Resumes ({len(resumes_tab4)})", expanded=False):
                st.write(f"**Total resumes to be screened: {len(resumes_tab4)}**")
                for i, resume in enumerate(resumes_tab4):
                    st.markdown(f"**{i+1}. {resume['filename']}** ({len(resume['content'])} characters)")
        else:
            st.error("❌ Please upload at least one resume in the 'Resumes' tab.")
        
        # Pre-screening validation
        if can_start:
            st.success("✅ All requirements met - Ready to start screening!")
        else:
            st.warning("⚠️ Please complete all required steps before starting the screening process.")
        
        # Start screening button with consistent design
        if st.button(
            '� Start AI Resume Screening',
            disabled=not can_start or st.session_state.get('running', False),
            type="primary",
            use_container_width=True,
            key="start_screening_button_tab4"
        ):
            if can_start:
                # Check usage limits before proceeding
                if not add_usage_monitoring():
                    st.stop()
                
                st.session_state.running = True
                
                job_profile_to_screen = st.session_state.get('job_profile', "")
                resumes_to_screen = st.session_state.get('resumes', [])

                with st.container():
                    st.markdown("---")
                    st.markdown("### 📊 Screening Progress")
                    
                    try:
                        screening_results = asyncio.run(
                            main_screening(job_profile_to_screen, resumes_to_screen, st.session_state.get('num_agents', 4))
                        )
                        st.session_state.screening_results = screening_results
                        st.session_state.running = False
                        
                        # Record usage after successful processing
                        tracker = UsageTracker()
                        total_cost = len(resumes_to_screen) * tracker.cost_per_resume
                        usage = tracker.record_usage(num_resumes=len(resumes_to_screen), actual_cost=total_cost)
                        
                        st.success(f"✅ Screening completed! Daily usage: ${usage['daily_cost']:.2f}")
                        
                        # Display results
                        display_screening_results(screening_results, resumes_to_screen)
                        
                    except Exception as e:
                        st.error(f"❌ An error occurred during screening: {str(e)}")
                        st.session_state.running = False
    
    # Display previous results if available
    if st.session_state.get('screening_results') and not st.session_state.get('running', False):
        st.markdown("---")
        st.markdown("### 📊 Latest Screening Results")
        display_screening_results(st.session_state.screening_results, st.session_state.get('resumes', []))

def display_screening_results(results, resumes):
    """Display the screening results in an organized format"""
    st.markdown("### 🎯 Matching Results")
    
    # Use actual plugin results for accurate scoring instead of mock data
    from src.plugins.resume_screening import ResumeScreeningPlugin
    
    plugin = ResumeScreeningPlugin()
    job_profile = st.session_state.get('job_profile', '')
    
    actual_results = []
    for resume in resumes:
        try:
            # Get actual plugin analysis
            result_json = plugin.analyze_resume(
                resume_content=resume['content'],
                job_profile=job_profile,
                candidate_name=resume['filename']
            )
            
            import json
            result = json.loads(result_json)
            
            # Create structured result with actual scores
            actual_results.append({
                "filename": resume["filename"],
                "score": result['overall_score'],
                "explanation": result['explanation'],
                "skill_score": result['skill_score'],
                "experience_score": result['experience_score'],
                "education_score": result['education_score'],
                "role_match_score": result['role_match_score'],
                "extracted_skills": result['extracted_skills'],
                "extracted_experience": result['extracted_experience'],
                "extracted_education": result['extracted_education']
            })
        except Exception as e:
            # Fallback for error cases
            actual_results.append({
                "filename": resume["filename"],
                "score": 0,
                "explanation": f"Error analyzing resume: {str(e)}",
                "skill_score": 0,
                "experience_score": 0,
                "education_score": 0,
                "role_match_score": 0,
                "extracted_skills": [],
                "extracted_experience": [],
                "extracted_education": []
            })
    
    # Sort by actual score
    actual_results.sort(key=lambda x: x["score"], reverse=True)
    sample_results = actual_results  # Use actual results instead of mock
    
    for i, result in enumerate(sample_results):
        with st.expander(f"🏆 #{i+1} - {result['filename']} - Score: {result['score']:.1f}%", expanded=i < 3):
            col1, col2 = st.columns([1, 2])
            
            with col1:
                st.markdown(f'<div class="matching-score">{result["score"]:.1f}%</div>', unsafe_allow_html=True)
                
                # Score interpretation with stricter thresholds
                if result["score"] >= 75:
                    st.success("🎯 Excellent Match")
                elif result["score"] >= 60:
                    st.warning("⚡ Good Match")  
                elif result["score"] >= 45:
                    st.info("🔍 Moderate Match")
                elif result["score"] >= 30:
                    st.warning("⚠️ Weak Match")
                else:
                    st.error("❌ Poor Match")
                
                # Detailed score breakdown
                st.markdown("**Score Breakdown:**")
                st.metric("Skills", f"{result.get('skill_score', 0):.1f}%")
                st.metric("Experience", f"{result.get('experience_score', 0):.1f}%") 
                st.metric("Education", f"{result.get('education_score', 0):.1f}%")
                st.metric("Role Match", f"{result.get('role_match_score', 0):.1f}%")
            
            with col2:
                st.markdown("📝 **Analysis Summary:**")
                st.write(result["explanation"])
                
                # Show extracted data
                if result.get('extracted_skills'):
                    st.markdown("**🔧 Detected Skills:**")
                    st.write(", ".join(result['extracted_skills'][:10]))  # Show first 10 skills
                
                if result.get('extracted_experience'):
                    st.markdown("**📈 Experience Data:**") 
                    st.write(", ".join(result['extracted_experience']))
                    
                if result.get('extracted_education'):
                    st.markdown("**🎓 Education Data:**")
                    st.write(", ".join(result['extracted_education']))
    
    # Download report buttons
    col1, col2 = st.columns(2)
    
    with col1:
        if st.button("📥 Download CSV Report", type="secondary"):
            if 'screening_results' in st.session_state and st.session_state.screening_results:
                # Generate and download CSV report
                report_data = generate_detailed_report(st.session_state.screening_results)
                st.download_button(
                    label="📄 Download CSV Report",
                    data=report_data,
                    file_name=f"resume_screening_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                    mime="text/csv",
                    key="download_csv_report"
                )
            else:
                st.error("No screening results available. Please run the screening first.")
    
    with col2:
        if st.button("📄 Download Summary Report", type="secondary"):
            if 'screening_results' in st.session_state and st.session_state.screening_results:
                # Generate and download summary text report
                summary_data = generate_summary_report(st.session_state.screening_results)
                st.download_button(
                    label="📄 Download Summary",
                    data=summary_data,
                    file_name=f"resume_screening_summary_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
                    mime="text/plain",
                    key="download_summary_report"
                )
            else:
                st.error("No screening results available. Please run the screening first.")

def generate_summary_report(results):
    """Generate a human-readable summary report"""
    import io
    output = io.StringIO()
    
    # Report header
    output.write("AI RESUME SCREENING SUMMARY REPORT\n")
    output.write("=" * 50 + "\n")
    output.write(f"Generated on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    output.write(f"Total Candidates Analyzed: {len(results)}\n\n")
    
    # Summary statistics
    scores = [result.get('overall_score', 0) for result in results]
    avg_score = sum(scores) / len(scores) if scores else 0
    
    highly_recommended = len([s for s in scores if s >= 75])
    recommended = len([s for s in scores if 60 <= s < 75])
    consider_caution = len([s for s in scores if 45 <= s < 60])
    not_recommended = len([s for s in scores if s < 45])
    
    output.write("SUMMARY STATISTICS:\n")
    output.write("-" * 20 + "\n")
    output.write(f"Average Score: {avg_score:.1f}%\n")
    output.write(f"Highly Recommended: {highly_recommended} candidates\n")
    output.write(f"Recommended: {recommended} candidates\n")
    output.write(f"Consider with Caution: {consider_caution} candidates\n")
    output.write(f"Not Recommended: {not_recommended} candidates\n\n")
    
    # Individual candidate details
    output.write("INDIVIDUAL CANDIDATE ANALYSIS:\n")
    output.write("=" * 50 + "\n\n")
    
    # Sort by score (highest first)
    sorted_results = sorted(results, key=lambda x: x.get('overall_score', 0), reverse=True)
    
    for i, result in enumerate(sorted_results, 1):
        output.write(f"RANK #{i}: {result.get('candidate', 'Unknown Candidate')}\n")
        output.write("-" * 30 + "\n")
        output.write(f"File: {result.get('filename', 'Unknown')}\n")
        output.write(f"Overall Score: {result.get('overall_score', 0):.1f}%\n")
        
        # Score breakdown
        output.write("\nScore Breakdown:\n")
        output.write(f"  • Skills Match: {result.get('skill_score', 0):.1f}%\n")
        output.write(f"  • Experience Match: {result.get('experience_score', 0):.1f}%\n")
        output.write(f"  • Education Match: {result.get('education_score', 0):.1f}%\n")
        output.write(f"  • Role Level Match: {result.get('role_match_score', 0):.1f}%\n")
        
        # Key findings
        skills = result.get('extracted_skills', [])[:3]
        if skills:
            output.write(f"\nKey Skills Found: {', '.join(skills)}\n")
        
        experience = result.get('extracted_experience', [])
        if experience:
            output.write(f"Experience: {', '.join(experience)} years\n")
            
        education = result.get('extracted_education', [])[:2]
        if education:
            output.write(f"Education: {', '.join(education)}\n")
        
        # Recommendation
        overall_score = result.get('overall_score', 0)
        if overall_score >= 75:
            recommendation = "HIGHLY RECOMMENDED ⭐⭐⭐"
        elif overall_score >= 60:
            recommendation = "RECOMMENDED ⭐⭐"
        elif overall_score >= 45:
            recommendation = "CONSIDER WITH CAUTION ⚠️"
        else:
            recommendation = "NOT RECOMMENDED ❌"
            
        output.write(f"\nRecommendation: {recommendation}\n")
        output.write("\n" + "="*50 + "\n\n")
    
    return output.getvalue()

def generate_detailed_report(results):
    """Generate a comprehensive CSV report of the screening results"""
    import io
    output = io.StringIO()
    
    # Create detailed CSV header
    output.write("Rank,Candidate_Name,Filename,Overall_Score,Skill_Score,Experience_Score,Education_Score,Role_Match_Score,Recommendation,Skills_Found,Experience_Years,Education_Background,Strengths,Areas_of_Concern\n")
    
    for i, result in enumerate(results):
        try:
            # Extract data safely with defaults
            overall_score = result.get('overall_score', 0)
            skill_score = result.get('skill_score', 0)
            experience_score = result.get('experience_score', 0) 
            education_score = result.get('education_score', 0)
            role_match_score = result.get('role_match_score', 0)
            
            # Determine recommendation based on score
            if overall_score >= 75:
                recommendation = "Highly Recommended"
            elif overall_score >= 60:
                recommendation = "Recommended" 
            elif overall_score >= 45:
                recommendation = "Consider with Caution"
            elif overall_score >= 30:
                recommendation = "Weak Match"
            else:
                recommendation = "Not Recommended"
            
            # Extract skills, experience, education
            skills = ", ".join(result.get('extracted_skills', [])[:5])  # Top 5 skills
            experience = ", ".join(result.get('extracted_experience', []))
            education = ", ".join(result.get('extracted_education', [])[:3])  # Top 3 education items
            
            # Extract strengths and concerns from explanation if available
            explanation = result.get('explanation', '')
            strengths = "Technical skills present" if skill_score > 50 else "Limited technical alignment"
            concerns = "Experience gaps" if experience_score < 50 else "Minor skill gaps"
            
            # Format CSV row with proper escaping
            candidate_name = result.get('candidate', f'Candidate_{i+1}')
            filename = result.get('filename', 'Unknown')
            
            output.write(f'{i+1},"{candidate_name}","{filename}",{overall_score:.1f}%,{skill_score:.1f}%,{experience_score:.1f}%,{education_score:.1f}%,{role_match_score:.1f}%,"{recommendation}","{skills}","{experience}","{education}","{strengths}","{concerns}"\n')
            
        except Exception as e:
            # Fallback for any parsing errors
            output.write(f'{i+1},"Parse Error","Error",0%,0%,0%,0%,0%,"Error","Error parsing result","","","Error","Error"\n')
    
    return output.getvalue()

def generate_report(results):
    """Legacy function - kept for compatibility"""
    return generate_detailed_report(results)

if __name__ == "__main__":
    app()